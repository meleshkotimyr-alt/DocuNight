import os
import sys
import json
from pathlib import Path

from PySide6.QtCore import Qt, QUrl, QTimer
from PySide6.QtGui import QAction, QKeySequence, QFont, QDesktopServices
from PySide6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout, QLabel,
    QPushButton, QFileDialog, QTabWidget, QTextBrowser, QLineEdit, QToolBar,
    QStatusBar, QMessageBox, QSplitter, QListWidget, QListWidgetItem, QFrame,
    QDialog, QDialogButtonBox, QComboBox, QSpinBox, QCheckBox, QStackedWidget
)

try:
    from PySide6.QtPdf import QPdfDocument
    from PySide6.QtPdfWidgets import QPdfView
    HAS_QTPDF = True
except Exception:
    HAS_QTPDF = False

try:
    from docx import Document
except Exception:
    Document = None

APP_NAME = "DocuNight"
APP_DIR = Path(os.environ.get("APPDATA", str(Path.home()))) / APP_NAME
APP_DIR.mkdir(parents=True, exist_ok=True)
STATE_FILE = APP_DIR / "state.json"

LIGHT_QSS = """
QMainWindow, QWidget { background: #f5f6f8; color: #202124; font-family: 'Segoe UI'; font-size: 10pt; }
QToolBar { background: #ffffff; border: 1px solid #e1e4e8; spacing: 7px; padding: 7px; }
QToolButton { background: transparent; color: #202124; padding: 7px 10px; border-radius: 7px; }
QToolButton:hover { background: #edf0f4; }
QFrame#sidebar { background: #ffffff; border: 1px solid #e1e4e8; border-radius: 12px; }
QLabel#brand { font-size: 20px; font-weight: 700; }
QLabel#section { color: #6d7480; font-size: 9pt; font-weight: 700; }
QLabel#muted { color: #68707c; }
QLabel#welcomeTitle { font-size: 24px; font-weight: 700; }
QPushButton { background: #f0f2f5; color: #202124; border: 1px solid #dfe3e8; border-radius: 8px; padding: 9px 11px; text-align: left; }
QPushButton:hover { background: #e6eaf0; }
QPushButton#primary { background: #2563eb; color: white; border: 0; padding: 11px 18px; text-align: center; }
QLineEdit, QSpinBox, QComboBox { background: #ffffff; color: #202124; border: 1px solid #d6dbe2; border-radius: 7px; padding: 8px; }
QTabWidget::pane { border: 1px solid #e1e4e8; border-radius: 9px; background: #ffffff; }
QTabBar::tab { background: #e9edf2; color: #202124; padding: 9px 12px; margin-right: 3px; border-top-left-radius: 6px; border-top-right-radius: 6px; }
QTabBar::tab:selected { background: #2563eb; color: white; }
QListWidget { background: #ffffff; color: #202124; border: 1px solid #e1e4e8; border-radius: 8px; padding: 4px; }
QListWidget::item { padding: 8px 5px; border-radius: 5px; }
QListWidget::item:hover { background: #edf0f4; }
QStatusBar { background: #ffffff; color: #626b78; }
QTextBrowser { background: #ffffff; color: #202124; border: 0; padding: 22px; }
"""

DARK_QSS = """
QMainWindow, QWidget { background: #17191f; color: #edf0f6; font-family: 'Segoe UI'; font-size: 10pt; }
QToolBar { background: #20232b; border: 1px solid #383e4b; spacing: 7px; padding: 7px; }
QToolButton { background: transparent; color: #edf0f6; padding: 7px 10px; border-radius: 7px; }
QToolButton:hover { background: #303646; }
QFrame#sidebar { background: #20232b; border: 1px solid #383e4b; border-radius: 12px; }
QLabel#brand { font-size: 20px; font-weight: 700; }
QLabel#section { color: #a4adbd; font-size: 9pt; font-weight: 700; }
QLabel#muted { color: #a4adbd; }
QLabel#welcomeTitle { font-size: 24px; font-weight: 700; }
QPushButton { background: #282c36; color: #edf0f6; border: 1px solid #383e4b; border-radius: 8px; padding: 9px 11px; text-align: left; }
QPushButton:hover { background: #303646; }
QPushButton#primary { background: #5268db; color: white; border: 0; padding: 11px 18px; text-align: center; }
QLineEdit, QSpinBox, QComboBox { background: #14161b; color: #edf0f6; border: 1px solid #383e4b; border-radius: 7px; padding: 8px; }
QTabWidget::pane { border: 1px solid #383e4b; border-radius: 9px; background: #20232b; }
QTabBar::tab { background: #282c36; color: #edf0f6; padding: 9px 12px; margin-right: 3px; border-top-left-radius: 6px; border-top-right-radius: 6px; }
QTabBar::tab:selected { background: #5268db; color: white; }
QListWidget { background: #14161b; color: #edf0f6; border: 1px solid #383e4b; border-radius: 8px; padding: 4px; }
QListWidget::item { padding: 8px 5px; border-radius: 5px; }
QListWidget::item:hover { background: #303646; }
QStatusBar { background: #20232b; color: #a4adbd; }
"""

def load_state():
    try:
        return json.loads(STATE_FILE.read_text(encoding="utf-8"))
    except Exception:
        return {"recent": [], "bookmarks": {}, "last_pages": {}, "theme": "light"}

def save_state(state):
    try:
        STATE_FILE.write_text(json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8")
    except Exception:
        pass

class DocumentTab(QWidget):
    def __init__(self, path, state, theme, on_page_changed=None):
        super().__init__()
        self.path = os.path.abspath(path)
        self.state = state
        self.theme = theme
        self.on_page_changed = on_page_changed
        self.pdf_document = None
        self.pdf_view = None
        self.browser = None
        self.page_texts = []
        self.current_page = 0
        self.ext = Path(path).suffix.lower()

        outer = QVBoxLayout(self)
        outer.setContentsMargins(0, 0, 0, 0)
        self.content = QWidget()
        self.content_layout = QVBoxLayout(self.content)
        self.content_layout.setContentsMargins(0, 0, 0, 0)
        outer.addWidget(self.content)

        if self.ext == ".pdf":
            self.load_pdf()
        elif self.ext == ".docx":
            self.load_docx()
        elif self.ext == ".doc":
            self.message("Старый формат Word (.doc)", "Откройте файл в Word или LibreOffice и сохраните его как .docx.")
        else:
            self.message("Формат не поддерживается", "Поддерживаются PDF и DOCX.")

        self.apply_theme(theme)

    def message(self, title, body):
        label = QLabel(f"<h2>{title}</h2><p>{body}</p>")
        label.setWordWrap(True)
        label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        label.setContentsMargins(28, 28, 28, 28)
        self.content_layout.addWidget(label)

    def load_pdf(self):
        if not HAS_QTPDF:
            self.message("Просмотр PDF недоступен", "Установите зависимости командой: py -m pip install -r requirements.txt")
            return
        self.pdf_document = QPdfDocument(self)
        status = self.pdf_document.load(self.path)
        if status != QPdfDocument.Status.Ready:
            # Loading can be asynchronous in some Qt versions; the view still receives the document.
            pass
        self.pdf_view = QPdfView()
        self.pdf_view.setDocument(self.pdf_document)
        self.pdf_view.setPageMode(QPdfView.PageMode.MultiPage)
        self.pdf_view.setZoomMode(QPdfView.ZoomMode.FitToWidth)
        self.content_layout.addWidget(self.pdf_view)

        try:
            count = self.pdf_document.pageCount()
            self.page_texts = [self.pdf_document.getAllText(i).text() for i in range(count)]
        except Exception:
            self.page_texts = []
        self.restore_page()

    def load_docx(self):
        if Document is None:
            self.message("Модуль Word не установлен", "Выполните: py -m pip install python-docx")
            return
        try:
            doc = Document(self.path)
            parts = []
            for p in doc.paragraphs:
                text = p.text
                if not text.strip():
                    continue
                escaped = (text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))
                style = (p.style.name or "").lower()
                if "title" in style:
                    parts.append(f"<h1>{escaped}</h1>")
                elif "heading 1" in style:
                    parts.append(f"<h2>{escaped}</h2>")
                elif "heading 2" in style:
                    parts.append(f"<h3>{escaped}</h3>")
                else:
                    parts.append(f"<p>{escaped}</p>")
            for table in doc.tables:
                parts.append("<table border='1' cellspacing='0' cellpadding='5'>")
                for row in table.rows:
                    parts.append("<tr>" + "".join(
                        "<td>" + cell.text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\\n", "<br>") + "</td>"
                        for cell in row.cells
                    ) + "</tr>")
                parts.append("</table><br>")
            self.browser = QTextBrowser()
            self.browser.setOpenExternalLinks(True)
            self.browser.setHtml("".join(parts) or "<p>В документе нет текста.</p>")
            self.content_layout.addWidget(self.browser)
        except Exception as e:
            self.message("Не удалось открыть DOCX", str(e))

    def restore_page(self):
        if not self.pdf_view or not self.pdf_document:
            return
        saved = int(self.state.get("last_pages", {}).get(self.path, 0) or 0)
        count = self.pdf_document.pageCount()
        if count:
            saved = max(0, min(saved, count - 1))
            QTimer.singleShot(250, lambda: self.go_to_page(saved, persist=False))

    def go_to_page(self, page, persist=True):
        if not self.pdf_view or not self.pdf_document:
            return False
        count = self.pdf_document.pageCount()
        if count <= 0:
            return False
        page = max(0, min(int(page), count - 1))
        self.current_page = page
        try:
            nav = self.pdf_view.pageNavigator()
            nav.jump(page, 0)
        except Exception:
            try:
                self.pdf_view.setPageMode(QPdfView.PageMode.SinglePage)
                self.pdf_view.pageNavigator().jump(page, 0)
            except Exception:
                return False
        if persist:
            self.state.setdefault("last_pages", {})[self.path] = page
            save_state(self.state)
        if self.on_page_changed:
            self.on_page_changed(self, page)
        return True

    def current_page_number(self):
        try:
            return int(self.pdf_view.pageNavigator().currentPage()) + 1
        except Exception:
            return self.current_page + 1

    def search(self, query):
        query = query.casefold().strip()
        if not query:
            return []
        if self.ext == ".pdf":
            results = []
            for i, text in enumerate(self.page_texts):
                if query in (text or "").casefold():
                    results.append((i, (text or "").casefold().count(query)))
            return results
        if self.browser:
            # Search DOCX rendered text using QTextBrowser's built-in find.
            found = self.browser.find(query, QTextBrowser.FindFlag(0))
            if not found:
                cursor = self.browser.textCursor()
                cursor.movePosition(cursor.MoveOperation.Start)
                self.browser.setTextCursor(cursor)
                found = self.browser.find(query)
            return [(0, 1)] if found else []
        return []

    def apply_theme(self, theme):
        self.theme = theme
        if self.browser:
            if theme == "dark":
                self.browser.setStyleSheet("QTextBrowser { background:#20232b; color:#edf0f6; border:0; padding:22px; }")
            else:
                self.browser.setStyleSheet("QTextBrowser { background:#ffffff; color:#202124; border:0; padding:22px; }")

class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.state = load_state()
        self.theme = self.state.get("theme", "light")
        self.setWindowTitle("DocuNight")
        self.resize(1280, 840)
        self.setMinimumSize(900, 600)
        self.setAcceptDrops(True)
        self.build_ui()
        self.apply_theme()

        # Restore last-used documents only if they still exist.
        for p in self.state.get("recent", [])[:5]:
            if os.path.isfile(p):
                self.add_document(p, record=False)

    def build_ui(self):
        toolbar = QToolBar("Главная")
        toolbar.setMovable(False)
        self.addToolBar(toolbar)

        act_open = QAction("📂  Открыть", self)
        act_open.setShortcut(QKeySequence.StandardKey.Open)
        act_open.triggered.connect(self.open_files)
        toolbar.addAction(act_open)

        act_search = QAction("⌕  Поиск", self)
        act_search.setShortcut(QKeySequence.StandardKey.Find)
        act_search.triggered.connect(self.focus_search)
        toolbar.addAction(act_search)

        self.search_input = QLineEdit()
        self.search_input.setPlaceholderText("Найти текст в открытом документе…")
        self.search_input.setMaximumWidth(300)
        self.search_input.returnPressed.connect(self.run_search)
        toolbar.addWidget(self.search_input)

        self.theme_action = QAction("☾  Тёмная тема", self)
        self.theme_action.triggered.connect(self.toggle_theme)
        toolbar.addAction(self.theme_action)

        central = QWidget()
        root = QHBoxLayout(central)
        root.setContentsMargins(12, 12, 12, 12)
        root.setSpacing(12)

        self.sidebar = QFrame()
        self.sidebar.setObjectName("sidebar")
        self.sidebar.setFixedWidth(250)
        side = QVBoxLayout(self.sidebar)
        side.setContentsMargins(12, 15, 12, 12)
        side.setSpacing(9)

        brand = QLabel("◈  DocuNight")
        brand.setObjectName("brand")
        side.addWidget(brand)
        subtitle = QLabel("Пространство для документов")
        subtitle.setObjectName("muted")
        side.addWidget(subtitle)
        side.addSpacing(12)

        section = QLabel("БЫСТРЫЙ ДОСТУП")
        section.setObjectName("section")
        side.addWidget(section)

        pdf_btn = QPushButton("📕  Открыть PDF")
        pdf_btn.clicked.connect(lambda: self.open_files("PDF (*.pdf)"))
        side.addWidget(pdf_btn)
        word_btn = QPushButton("📝  Открыть Word")
        word_btn.clicked.connect(lambda: self.open_files("Word (*.docx *.doc)"))
        side.addWidget(word_btn)

        self.bookmark_button = QPushButton("🔖  Добавить закладку")
        self.bookmark_button.clicked.connect(self.add_bookmark)
        side.addWidget(self.bookmark_button)

        side.addSpacing(12)
        recent_label = QLabel("НЕДАВНИЕ ФАЙЛЫ")
        recent_label.setObjectName("section")
        side.addWidget(recent_label)
        self.recent_list = QListWidget()
        self.recent_list.itemDoubleClicked.connect(self.open_recent)
        side.addWidget(self.recent_list, 1)

        bookmarks_label = QLabel("ЗАКЛАДКИ")
        bookmarks_label.setObjectName("section")
        side.addWidget(bookmarks_label)
        self.bookmark_list = QListWidget()
        self.bookmark_list.itemDoubleClicked.connect(self.open_bookmark)
        side.addWidget(self.bookmark_list, 1)

        hint = QLabel("Перетащи PDF или DOCX в окно.\nCtrl+O — открыть · Ctrl+F — поиск")
        hint.setObjectName("muted")
        hint.setWordWrap(True)
        side.addWidget(hint)

        self.tabs = QTabWidget()
        self.tabs.setDocumentMode(True)
        self.tabs.setTabsClosable(True)
        self.tabs.tabCloseRequested.connect(self.close_tab)
        self.tabs.currentChanged.connect(self.update_status)
        self.show_welcome()

        splitter = QSplitter(Qt.Orientation.Horizontal)
        splitter.addWidget(self.sidebar)
        splitter.addWidget(self.tabs)
        splitter.setStretchFactor(0, 0)
        splitter.setStretchFactor(1, 1)
        root.addWidget(splitter)
        self.setCentralWidget(central)
        self.setStatusBar(QStatusBar())
        self.refresh_recent()
        self.refresh_bookmarks()
        self.statusBar().showMessage("Готов к работе")

    def show_welcome(self):
        welcome = QWidget()
        layout = QVBoxLayout(welcome)
        layout.setAlignment(Qt.AlignmentFlag.AlignCenter)
        title = QLabel("Все документы в одном месте")
        title.setObjectName("welcomeTitle")
        title.setAlignment(Qt.AlignmentFlag.AlignCenter)
        layout.addWidget(title)
        desc = QLabel("Откройте PDF или Word, чтобы начать.\nМожно перетащить файл прямо в это окно.")
        desc.setObjectName("muted")
        desc.setAlignment(Qt.AlignmentFlag.AlignCenter)
        layout.addWidget(desc)
        layout.addSpacing(18)
        btn = QPushButton("📂  Выбрать документ")
        btn.setObjectName("primary")
        btn.clicked.connect(self.open_files)
        btn.setMinimumWidth(190)
        btn.setMinimumHeight(42)
        layout.addWidget(btn, alignment=Qt.AlignmentFlag.AlignCenter)
        self.tabs.addTab(welcome, "Главная")

    def open_files(self, filter_text=None):
        filt = (filter_text + ";;Все поддерживаемые документы (*.pdf *.docx *.doc);;Все файлы (*)") if filter_text else "Документы (*.pdf *.docx *.doc);;PDF (*.pdf);;Word (*.docx *.doc);;Все файлы (*)"
        paths, _ = QFileDialog.getOpenFileNames(self, "Открыть документы", "", filt)
        for p in paths:
            self.add_document(p)

    def add_document(self, path, record=True):
        if not path or not os.path.isfile(path):
            return
        ext = Path(path).suffix.lower()
        if ext not in (".pdf", ".docx", ".doc"):
            self.statusBar().showMessage("Поддерживаются PDF и Word DOCX.", 4000)
            return
        path = os.path.abspath(path)
        for i in range(self.tabs.count()):
            tab = self.tabs.widget(i)
            if isinstance(tab, DocumentTab) and tab.path == path:
                self.tabs.setCurrentIndex(i)
                return
        if self.tabs.count() == 1 and not isinstance(self.tabs.widget(0), DocumentTab):
            self.tabs.removeTab(0)
        tab = DocumentTab(path, self.state, self.theme, self.page_changed)
        icon = "📕" if ext == ".pdf" else "📝"
        idx = self.tabs.addTab(tab, f"{icon} {Path(path).name}")
        self.tabs.setTabToolTip(idx, path)
        self.tabs.setCurrentIndex(idx)
        if record:
            recent = self.state.setdefault("recent", [])
            self.state["recent"] = [path] + [p for p in recent if p != path]
            self.state["recent"] = self.state["recent"][:15]
            save_state(self.state)
            self.refresh_recent()
        self.statusBar().showMessage(f"Открыт: {Path(path).name}", 3500)

    def page_changed(self, tab, page):
        self.statusBar().showMessage(f"{Path(tab.path).name} — страница {page + 1}", 2500)

    def current_doc(self):
        tab = self.tabs.currentWidget()
        return tab if isinstance(tab, DocumentTab) else None

    def focus_search(self):
        self.search_input.setFocus()
        self.search_input.selectAll()

    def run_search(self):
        query = self.search_input.text().strip()
        tab = self.current_doc()
        if not tab:
            self.statusBar().showMessage("Сначала откройте документ.", 3000)
            return
        if not query:
            return
        results = tab.search(query)
        if tab.ext == ".pdf":
            if not results:
                QMessageBox.information(self, "Поиск", "Совпадений не найдено. Некоторые PDF состоят из сканов и не содержат распознаваемого текста.")
                return
            dlg = QDialog(self)
            dlg.setWindowTitle("Результаты поиска")
            dlg.resize(500, 350)
            layout = QVBoxLayout(dlg)
            label = QLabel(f"Найдено страниц с совпадениями: {len(results)}")
            layout.addWidget(label)
            listing = QListWidget()
            for page, count in results:
                item = QListWidgetItem(f"Страница {page + 1}  ·  совпадений: {count}")
                item.setData(Qt.ItemDataRole.UserRole, page)
                listing.addItem(item)
            layout.addWidget(listing)
            buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Close)
            buttons.rejected.connect(dlg.reject)
            buttons.accepted.connect(dlg.accept)
            layout.addWidget(buttons)
            def jump(_item):
                page = _item.data(Qt.ItemDataRole.UserRole)
                tab.go_to_page(page)
                dlg.accept()
            listing.itemDoubleClicked.connect(jump)
            if results:
                listing.setCurrentRow(0)
            dlg.exec()
        else:
            QMessageBox.information(self, "Поиск", "Текст найден." if results else "Совпадений не найдено.")

    def add_bookmark(self):
        tab = self.current_doc()
        if not tab:
            self.statusBar().showMessage("Сначала откройте документ.", 3000)
            return
        page = tab.current_page_number() - 1 if tab.ext == ".pdf" else 0
        entry = {"path": tab.path, "page": max(0, page), "label": f"{Path(tab.path).name} — страница {page + 1}"}
        bookmarks = self.state.setdefault("bookmarks", [])
        key = (entry["path"], entry["page"])
        if not any((b.get("path"), int(b.get("page", 0))) == key for b in bookmarks):
            bookmarks.append(entry)
            save_state(self.state)
            self.refresh_bookmarks()
            self.statusBar().showMessage("Закладка сохранена.", 2500)
        else:
            self.statusBar().showMessage("Для этой страницы закладка уже существует.", 2500)

    def refresh_bookmarks(self):
        self.bookmark_list.clear()
        for b in self.state.get("bookmarks", []):
            item = QListWidgetItem("🔖  " + b.get("label", Path(b.get("path", "")).name))
            item.setData(Qt.ItemDataRole.UserRole, b)
            item.setToolTip(b.get("path", ""))
            self.bookmark_list.addItem(item)

    def open_bookmark(self, item):
        b = item.data(Qt.ItemDataRole.UserRole)
        path = b.get("path", "")
        if not os.path.isfile(path):
            QMessageBox.warning(self, "Закладка", "Исходный файл не найден. Возможно, он был перемещён или удалён.")
            return
        self.add_document(path)
        tab = self.current_doc()
        if tab and tab.ext == ".pdf":
            tab.go_to_page(int(b.get("page", 0)))

    def refresh_recent(self):
        self.recent_list.clear()
        for p in self.state.get("recent", [])[:15]:
            if os.path.isfile(p):
                icon = "📕" if Path(p).suffix.lower() == ".pdf" else "📝"
                item = QListWidgetItem(f"{icon}  {Path(p).name}")
                item.setData(Qt.ItemDataRole.UserRole, p)
                item.setToolTip(p)
                self.recent_list.addItem(item)

    def open_recent(self, item):
        self.add_document(item.data(Qt.ItemDataRole.UserRole))

    def close_tab(self, index):
        self.tabs.removeTab(index)
        if self.tabs.count() == 0:
            self.show_welcome()

    def toggle_theme(self):
        self.theme = "dark" if self.theme == "light" else "light"
        self.state["theme"] = self.theme
        save_state(self.state)
        self.apply_theme()
        for i in range(self.tabs.count()):
            tab = self.tabs.widget(i)
            if isinstance(tab, DocumentTab):
                tab.apply_theme(self.theme)

    def apply_theme(self):
        self.setStyleSheet(DARK_QSS if self.theme == "dark" else LIGHT_QSS)
        self.theme_action.setText("☀  Светлая тема" if self.theme == "dark" else "☾  Тёмная тема")

    def dragEnterEvent(self, event):
        if event.mimeData().hasUrls():
            event.acceptProposedAction()
        else:
            event.ignore()

    def dropEvent(self, event):
        for url in event.mimeData().urls():
            if url.isLocalFile():
                self.add_document(url.toLocalFile())
        event.acceptProposedAction()

    def update_status(self, *_):
        tab = self.current_doc()
        if tab:
            self.statusBar().showMessage(tab.path)
        else:
            self.statusBar().showMessage("Готов к работе")

def main():
    app = QApplication(sys.argv)
    app.setApplicationName(APP_NAME)
    app.setFont(QFont("Segoe UI", 10))
    win = MainWindow()
    win.show()
    sys.exit(app.exec())

if __name__ == "__main__":
    main()
